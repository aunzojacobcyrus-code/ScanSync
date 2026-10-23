"""
Spreadsheet export for attendance records.

* **XLSX** — a *Summary* sheet (one row per class/date with totals) followed by
  one worksheet per class + date group, each listing every student's name, ID,
  date, time in and status. Status cells are colour-coded.
* **CSV** — a single flat file sorted (grouped) by class then date, so it opens
  cleanly in Excel / Google Sheets and can be pivoted.
"""
import csv
import re
from collections import OrderedDict
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from .models import AttendanceRecord

COLUMNS = ["Class Code", "Subject", "Section", "Date", "Student ID", "Last Name", "First Name",
           "Middle Name", "Time In", "Status", "Method", "Note"]


def ordered_records(records):
    return records.select_related("session__school_class__section", "student").order_by(
        "session__school_class__code", "session__school_class__section__name",
        "session__date", "student__last_name", "student__first_name",
    )


def record_row(r):
    cls = r.session.school_class
    t = timezone.localtime(r.time_in).strftime("%I:%M:%S %p") if r.time_in else ""
    return [cls.code, cls.name, cls.section.name, r.session.date.isoformat(), r.student.student_id,
            r.student.last_name, r.student.first_name, r.student.middle_name, t,
            r.get_status_display(), r.get_method_display(), r.note]


def grouped(records):
    """OrderedDict[(session) -> [records]] preserving class/date ordering."""
    groups = OrderedDict()
    for r in ordered_records(records):
        groups.setdefault(r.session, []).append(r)
    return groups


def export_filename(ext, scope="attendance"):
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    return f"ScanSync_{scope}_{stamp}.{ext}"


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
def export_csv(records, scope="attendance"):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{export_filename("csv", scope)}"'
    response.write("\ufeff")  # BOM so Excel opens UTF-8 names (ñ, é) correctly
    writer = csv.writer(response)
    writer.writerow(COLUMNS)
    for r in ordered_records(records):
        writer.writerow(record_row(r))
    return response


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #
def _sheet_title(session, used):
    cls = session.school_class
    base = re.sub(r"[\[\]\*\?/\\:]", "", f"{cls.code} {cls.section.name} {session.date:%m-%d}")[:31]
    title, n = base, 2
    while title in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title)
    return title


def export_xlsx(records, scope="attendance"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    brand = "4F46E5"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=brand)
    title_font = Font(bold=True, size=14, color="1E1B4B")
    muted = Font(color="64748B", size=10)
    thin = Side(style="thin", color="E2E8F0")
    border = Border(bottom=thin)
    status_fill = {
        "Present": PatternFill("solid", fgColor="D1FAE5"),
        "Late": PatternFill("solid", fgColor="FEF3C7"),
        "Absent": PatternFill("solid", fgColor="FFE4E6"),
    }

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 22

    def autosize(ws, min_w=10, max_w=40):
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))

    groups = grouped(records)
    wb = Workbook()

    # ---- Summary ---------------------------------------------------------- #
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = "ScanSync Attendance Report", title_font
    ws["A2"] = f"Generated {timezone.localtime():%B %d, %Y %I:%M %p} · {len(groups)} class session(s)"
    ws["A2"].font = muted
    head = ["Class Code", "Subject", "Section", "Date", "Start", "Grace (min)",
            "Present", "Late", "Absent", "Total", "Attendance %"]
    ws.append([])
    ws.append(head)
    style_header(ws, 4, len(head))
    for session, rows in groups.items():
        cls = session.school_class
        counts = {s: 0 for s in ("Present", "Late", "Absent")}
        for r in rows:
            counts[r.get_status_display()] += 1
        total = len(rows)
        rate = (counts["Present"] + counts["Late"]) / total if total else 0
        ws.append([cls.code, cls.name, cls.section.name, session.date, session.start_time.strftime("%I:%M %p"),
                   session.grace_period_minutes, counts["Present"], counts["Late"], counts["Absent"], total, rate])
        ws.cell(row=ws.max_row, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=ws.max_row, column=11).number_format = "0%"
    ws.freeze_panes = "A5"
    autosize(ws)
    ws.column_dimensions["A"].width = 14

    # ---- One sheet per class + date -------------------------------------- #
    used = {"Summary"}
    detail_head = ["#", "Student ID", "Last Name", "First Name", "Middle Name", "Date", "Time In", "Status", "Method", "Note"]
    for session, rows in groups.items():
        cls = session.school_class
        sh = wb.create_sheet(_sheet_title(session, used))
        sh["A1"], sh["A1"].font = f"{cls.code} – {cls.name}", title_font
        sh["A2"] = (f"Section {cls.section.name} · {session.date:%A, %B %d, %Y} · "
                    f"Start {session.start_time:%I:%M %p} · Grace {session.grace_period_minutes} min")
        sh["A2"].font = muted
        sh.append([])
        sh.append(detail_head)
        style_header(sh, 4, len(detail_head))
        for i, r in enumerate(rows, 1):
            t = timezone.localtime(r.time_in).strftime("%I:%M:%S %p") if r.time_in else ""
            status = r.get_status_display()
            sh.append([i, r.student.student_id, r.student.last_name, r.student.first_name, r.student.middle_name,
                       session.date, t, status, r.get_method_display(), r.note])
            row = sh.max_row
            sh.cell(row=row, column=6).number_format = "yyyy-mm-dd"
            sh.cell(row=row, column=8).fill = status_fill.get(status)
            sh.cell(row=row, column=8).font = Font(bold=True)
            for c in range(1, len(detail_head) + 1):
                sh.cell(row=row, column=c).border = border
        sh.freeze_panes = "A5"
        sh.auto_filter.ref = f"A4:{get_column_letter(len(detail_head))}{max(sh.max_row, 4)}"
        autosize(sh)
        sh.column_dimensions["A"].width = 6

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{export_filename("xlsx", scope)}"'
    return response


def export(records, fmt, scope="attendance"):
    return export_xlsx(records, scope) if fmt == "xlsx" else export_csv(records, scope)


__all__ = ["export", "export_csv", "export_xlsx", "AttendanceRecord"]
